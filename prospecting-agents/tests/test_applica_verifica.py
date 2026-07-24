"""Test di scripts/applica_verifica: blocco RISCHIO ALTO, correzione CORREGGERE,
esclusione a valle nel queue builder."""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from scripts import applica_verifica as av  # noqa: E402
from queue import queue_builder as qb  # noqa: E402

REPORT_COLS = ["sheet", "nome", "azienda", "ruolo", "email", "verification",
               "Pattern", "Confidenza", "MX", "SMTP", "Punteggio", "Indirizzo suggerito"]


def _make_report(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verifica"
    ws.append(REPORT_COLS)
    # RISCHIO ALTO → blocca
    ws.append(["Nuovi contatti", "A Uno", "Ente1", "Manager", "a.uno@dead.zz",
               "valid", "", "", "no", "", "RISCHIO ALTO", ""])
    # CORREGGERE conf 0.99 → corregge
    ws.append(["Nuovi contatti", "B Due", "Ente2", "Manager", "b.due@ente2.it",
               "accept_all", "nome.cognome", 0.99, "sì", "", "CORREGGERE",
               "bruno.due@ente2.it"])
    # CORREGGERE conf 0.64 → sotto soglia, salta
    ws.append(["Nuovi contatti", "C Tre", "Ente3", "Manager", "c.tre@ente3.it",
               "accept_all", "nome.cognome", 0.64, "sì", "", "CORREGGERE",
               "carla.tre@ente3.it"])
    # OK → ignorato
    ws.append(["Nuovi contatti", "D Quattro", "Ente4", "Manager", "d.q@ente4.it",
               "valid", "", "", "sì", "", "OK", ""])
    p = tmp_path / "report.xlsx"
    wb.save(p)
    return p


def _make_master(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    nc = wb.active
    nc.title = "Nuovi contatti"
    nc.append(["Score", "Azione", "Segmento", "Persona", "Nome", "Azienda", "Ruolo",
               "Dipartimento", "Città", "Email", "Verification", "Confidence", "Nota"])
    nc.append([100, "N", "PA", "Q", "A Uno", "Ente1", "Manager", "D", "Roma",
               "a.uno@dead.zz", "valid", 90, ""])
    nc.append([100, "N", "PA", "Q", "B Due", "Ente2", "Manager", "D", "Roma",
               "b.due@ente2.it", "accept_all", 80, ""])
    nc.append([100, "N", "PA", "Q", "C Tre", "Ente3", "Manager", "D", "Roma",
               "c.tre@ente3.it", "accept_all", 80, ""])
    p = tmp_path / "master.xlsx"
    wb.save(p)
    return p


def _cfg(tmp_path):
    return {"paths": {"backup": str(tmp_path / "b"), "logs": str(tmp_path / "l")}}


def test_blocca_rischio_alto(tmp_path):
    master, report = _make_master(tmp_path), _make_report(tmp_path)
    res = av.applica(master, report, _cfg(tmp_path), apply=True)
    assert res["blocked"] == 1
    nc = openpyxl.load_workbook(master)["Nuovi contatti"]
    nota = {nc.cell(r, 10).value: nc.cell(r, 13).value for r in range(2, nc.max_row + 1)}
    assert "Non inviare" in (nota["a.uno@dead.zz"] or "")


def test_corregge_solo_sopra_soglia(tmp_path):
    master, report = _make_master(tmp_path), _make_report(tmp_path)
    res = av.applica(master, report, _cfg(tmp_path), apply=True)
    assert res["corrected"] == 1
    nc = openpyxl.load_workbook(master)["Nuovi contatti"]
    emails = {nc.cell(r, 10).value for r in range(2, nc.max_row + 1)}
    assert "bruno.due@ente2.it" in emails          # corretto (conf 0.99)
    assert "b.due@ente2.it" not in emails
    assert "c.tre@ente3.it" in emails              # NON corretto (conf 0.64 < 0.8)
    # originale annotato nella Nota
    nota = {nc.cell(r, 10).value: nc.cell(r, 13).value for r in range(2, nc.max_row + 1)}
    assert "b.due@ente2.it" in (nota["bruno.due@ente2.it"] or "")


def test_dry_run_non_scrive(tmp_path):
    master, report = _make_master(tmp_path), _make_report(tmp_path)
    before = master.read_bytes()
    res = av.applica(master, report, _cfg(tmp_path), apply=False)
    assert res["blocked"] == 1 and res["corrected"] == 1
    assert master.read_bytes() == before


def test_idempotente(tmp_path):
    master, report = _make_master(tmp_path), _make_report(tmp_path)
    av.applica(master, report, _cfg(tmp_path), apply=True)
    res2 = av.applica(master, report, _cfg(tmp_path), apply=True)
    assert res2["blocked"] == 0                    # già bloccato
    # il CORREGGERE ora non trova più b.due (già rinominato) → salta, non ricorregge
    assert res2["corrected"] == 0


def test_queue_builder_esclude_i_bloccati(tmp_path):
    """Accettazione F2: dopo --apply il queue builder esclude i bloccati."""
    master, report = _make_master(tmp_path), _make_report(tmp_path)
    av.applica(master, report, _cfg(tmp_path), apply=True)
    wb = openpyxl.load_workbook(master, data_only=True)
    _, _, blocked = qb.build_exclusions(wb)
    assert "a.uno@dead.zz" in blocked
    # e non compare in un batch costruito da Nuovi contatti
    from datetime import datetime
    cfg = {"caps": {"totale": 30, "fredde": 12, "follow_up": 10, "riprese": 5,
                    "accept_all": 3, "stesso_dominio": 8},
           "domini_congelati": [], "soglia_scorta": 60,
           "sender_email": "io@me.it", "sender_name": "Io",
           "paths": {"bozze": str(tmp_path / "bz")}}
    batch = qb.select_batch(wb, cfg, today=datetime(2026, 7, 10))
    assert "a.uno@dead.zz" not in {b["email"].lower() for b in batch}
