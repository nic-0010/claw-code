"""Applica al master gli esiti del verificatore email (Fase 2.2).

Legge il report `reports/verifica_email_YYYYMMDD.xlsx` (prodotto da
verifier.email_verifier) e fa SOLO due cose, con backup + --apply:

  (a) RISCHIO ALTO → scrive "Non inviare — rischio alto" nella Nota del contatto
      (col M "Nota" in Nuovi contatti, col P "Note" nelle code). Il queue builder
      esclude da solo i contatti così marcati.
  (b) CORREGGERE con confidenza pattern ≥ soglia (default 0.8) → sostituisce
      l'indirizzo con "Indirizzo suggerito", annotando l'originale nella Nota.

Non tocca nient'altro. Dry-run di default: mostra l'elenco completo prima di
scrivere. Idempotente. Mai colonne-formula.

Uso:
    python -m scripts.applica_verifica --config config.yaml                 # dry-run
    python -m scripts.applica_verifica --config config.yaml --apply
    python -m scripts.applica_verifica --config config.yaml --report reports/verifica_email_20260716.xlsx --apply
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

BLOCK_NOTE = "Non inviare — rischio alto"
RISCHIO_ALTO = "RISCHIO ALTO"
CORREGGERE = "CORREGGERE"
DEFAULT_CONF = 0.8


def _latest_report(reports_dir: Path) -> Path | None:
    reports = sorted(reports_dir.glob("verifica_email_*.xlsx"))
    return reports[-1] if reports else None


def _read_report(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Verifica"] if "Verifica" in wb.sheetnames else wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        rows.append({headers[i]: r[i] for i in range(len(headers))})
    return rows


def _header_map(ws) -> dict[str, int]:
    return {str(ws.cell(1, c).value).strip(): c
            for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def _col(headers: dict[str, int], *needles: str) -> int | None:
    for name, idx in headers.items():
        if any(n in name.lower() for n in needles):
            return idx
    return None


def _to_float(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def applica(master_path, report_path, cfg, *, apply: bool,
            conf_soglia: float = DEFAULT_CONF):
    from common import io_master as io

    log = io.RunLog(component="applica_verifica", apply=apply)
    rows = _read_report(Path(report_path))
    wb = io.load(master_path)

    # cache header/col per foglio
    sheet_cols: dict[str, dict] = {}

    def cols_for(sheet: str):
        if sheet not in sheet_cols:
            ws = wb[sheet]
            hm = _header_map(ws)
            sheet_cols[sheet] = {
                "ws": ws, "email": _col(hm, "email"),
                "nota": _col(hm, "nota", "note"),
            }
        return sheet_cols[sheet]

    def find_row(ws, col_email: int, email: str) -> int | None:
        target = email.strip().lower()
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(r, col_email).value or "").strip().lower() == target:
                return r
        return None

    blocked = corrected = skipped = 0
    report_lines = []
    for row in rows:
        sheet = str(row.get("sheet") or "").strip()
        email = str(row.get("email") or "").strip()
        punteggio = str(row.get("Punteggio") or "").strip()
        if not sheet or not email or sheet not in wb.sheetnames:
            continue
        c = cols_for(sheet)
        if not c["email"] or not c["nota"]:
            continue
        r = find_row(c["ws"], c["email"], email)
        if r is None:
            continue

        if punteggio == RISCHIO_ALTO:
            nota_att = str(c["ws"].cell(r, c["nota"]).value or "")
            if BLOCK_NOTE.lower() in nota_att.lower():
                skipped += 1
                continue
            nuovo = (nota_att.rstrip() + " · " + BLOCK_NOTE).strip(" ·") \
                if nota_att.strip() else BLOCK_NOTE
            io.set_cell(c["ws"], r, c["nota"], nuovo, sheet_name=sheet, log=log)
            blocked += 1
            report_lines.append(f"BLOCCA   {sheet:22s} {email}")

        elif punteggio == CORREGGERE:
            suggerito = str(row.get("Indirizzo suggerito") or "").strip()
            conf = _to_float(row.get("Confidenza"))
            if not suggerito or conf < conf_soglia:
                skipped += 1
                continue
            if suggerito.lower() == email.lower():
                skipped += 1
                continue
            io.set_cell(c["ws"], r, c["email"], suggerito, sheet_name=sheet, log=log)
            nota_att = str(c["ws"].cell(r, c["nota"]).value or "")
            ann = f"Indirizzo corretto da {email} (conf {conf:.2f})"
            nuovo = (nota_att.rstrip() + " · " + ann).strip(" ·") \
                if nota_att.strip() else ann
            io.set_cell(c["ws"], r, c["nota"], nuovo, sheet_name=sheet, log=log)
            corrected += 1
            report_lines.append(f"CORREGGE {sheet:22s} {email} → {suggerito} (conf {conf:.2f})")

    io.save(wb, master_path, apply=apply,
            backup_dir=cfg.get("paths", {}).get("backup", "backup"), log=log)
    log.notes.append(f"bloccati {blocked} · corretti {corrected} · saltati {skipped}")
    log.save(cfg.get("paths", {}).get("logs", "logs"))
    return {"blocked": blocked, "corrected": corrected, "skipped": skipped,
            "lines": report_lines, "log": log}


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Applica gli esiti del verificatore (Fase 2.2)")
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--report", default=None, help="path del report (default: l'ultimo)")
    ap.add_argument("--conf", type=float, default=DEFAULT_CONF,
                    help="soglia confidenza pattern per i CORREGGERE (default 0.8)")
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args(argv)

    import yaml

    cfg = yaml.safe_load(open(args.config, encoding="utf-8"))
    reports_dir = Path(cfg.get("paths", {}).get("reports", "reports"))
    report = Path(args.report) if args.report else _latest_report(reports_dir)
    if not report or not report.exists():
        print(f"Report non trovato in {reports_dir} (esegui prima il verificatore).")
        return 1

    res = applica(cfg["master_path"], report, cfg, apply=args.apply, conf_soglia=args.conf)
    print(f"Report: {report}")
    for line in res["lines"]:
        print(f"  {line}")
    print(f"Totali: {res['blocked']} da bloccare, {res['corrected']} da correggere, "
          f"{res['skipped']} saltati")
    if not args.apply:
        print("DRY-RUN: nessuna scrittura (usa --apply). Backup automatico all'apply.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
